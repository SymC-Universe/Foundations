import json
from pathlib import Path

import numpy as np
from openpyxl import Workbook

from d02a_physical_v02 import (
    load_frozen_carrier_matrix,
    normalize_carrier_shape,
    pairwise_carrier_metrics,
)


def test_matrix_parser_uses_independent_q_energy_axes(tmp_path):
    path = tmp_path / "fig3.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Panel a"

    q = [0.5, 0.75, 1.0]
    energy = [0.0, 1.0, 2.0, 3.0]
    for i, value in enumerate(q, 3):
        ws[f"A{i}"] = value
    for i, value in enumerate(energy, 3):
        ws[f"B{i}"] = value

    matrix = np.arange(12, dtype=float).reshape(4, 3) + 1.0
    for i in range(4):
        for j in range(3):
            ws.cell(row=3 + i, column=3 + j, value=float(matrix[i, j]))
    wb.save(path)

    result = load_frozen_carrier_matrix(
        path,
        {
            "sheet": "Panel a",
            "q_axis_range": "A3:A5",
            "energy_axis_range": "B3:B6",
            "intensity_matrix_range": "C3:E6",
            "matrix_shape_energy_by_q": [4, 3],
        },
    )
    assert np.allclose(result["q"], q)
    assert np.allclose(result["energy"], energy)
    assert result["intensity"].shape == (4, 3)
    assert np.allclose(result["intensity"], matrix)


def test_per_q_normalization_is_scale_invariant():
    q = np.array([0.5, 1.0])
    energy = np.array([0.0, 1.0, 2.0])
    base = np.array([
        [1.0, 3.0],
        [2.0, 4.0],
        [1.0, 3.0],
    ])
    common = np.array([0.0, 1.0, 2.0])

    a = normalize_carrier_shape(q, energy, base, common)
    b = normalize_carrier_shape(q, energy, 17.0 * base, common)
    assert np.allclose(a["normalized_intensity"], b["normalized_intensity"])
    for x, y in zip(a["q_metrics"], b["q_metrics"]):
        assert x["normalized_energy_centroid_meV"] == y["normalized_energy_centroid_meV"]
        assert x["normalized_energy_rms_width_meV"] == y["normalized_energy_rms_width_meV"]
        assert x["peak_energy_meV_on_common_grid"] == y["peak_energy_meV_on_common_grid"]


def test_negative_measurement_noise_is_clipped_before_shape_normalization():
    q = np.array([0.5])
    energy = np.array([0.0, 1.0, 2.0])
    intensity = np.array([[-1.0], [2.0], [2.0]])
    common = energy.copy()
    result = normalize_carrier_shape(q, energy, intensity, common)
    weights = result["normalized_intensity"][:, 0]
    assert np.all(weights >= 0.0)
    assert np.isclose(weights.sum(), 1.0)
    assert np.isclose(weights[0], 0.0)


def test_pairwise_metrics_report_zero_difference_for_scaled_identical_maps():
    q = np.array([0.5, 0.75, 1.0])
    energy = np.array([0.0, 1.0, 2.0])
    intensity = np.array([
        [1.0, 3.0, 2.0],
        [2.0, 4.0, 4.0],
        [1.0, 3.0, 1.0],
    ])
    common = energy.copy()
    a_shape = normalize_carrier_shape(q, energy, intensity, common)
    b_shape = normalize_carrier_shape(q, energy, 9.0 * intensity, common)
    a = {"temperature_K": 300.0, "q_metrics": a_shape["q_metrics"]}
    b = {"temperature_K": 385.0, "q_metrics": b_shape["q_metrics"]}
    m = pairwise_carrier_metrics(a, b)
    assert np.isclose(m["centroid_rms_difference_meV"], 0.0)
    assert np.isclose(m["width_rms_difference_meV"], 0.0)
    assert np.isclose(m["peak_energy_rms_difference_meV"], 0.0)
