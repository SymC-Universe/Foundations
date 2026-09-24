import json
from pathlib import Path

from openpyxl import Workbook

from d02a_physical import (
    extract_linewidth_table,
    phase_for_temperature,
    run_d02a,
)


def test_phase_classifier_is_frozen_and_value_independent():
    assert phase_for_temperature(50) == "ORTHORHOMBIC"
    assert phase_for_temperature(359.999) == "ORTHORHOMBIC"
    assert phase_for_temperature(360) == "TETRAGONAL"
    assert phase_for_temperature(409.999) == "TETRAGONAL"
    assert phase_for_temperature(410) == "CUBIC"


def test_linewidth_parser_retains_all_complete_rows(tmp_path):
    path = tmp_path / "fig4.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "4b"
    ws["A1"] = "M-point linewidths"
    ws["A2"] = "temperature (K)"
    ws["B2"] = "Linewidth (meV)"
    ws["C2"] = "Linewidth error (meV)"
    rows = [
        (100, 1.0, 0.1),
        (300, 2.0, 0.2),
        (385, 3.0, 0.3),
        (500, 4.0, 0.4),
    ]
    for idx, row in enumerate(rows, 3):
        for col, value in zip("ABC", row):
            ws[f"{col}{idx}"] = value
    ws["A7"] = 600
    ws["B7"] = 5.0
    wb.save(path)

    parsed = extract_linewidth_table(path)
    assert len(parsed) == 4
    assert [r["temperature_K"] for r in parsed] == [100.0, 300.0, 385.0, 500.0]
    assert parsed[2]["phase"] == "TETRAGONAL"
    assert parsed[3]["phase"] == "CUBIC"


def test_frozen_contract_refuses_chi_even_when_physical_source_is_processed(tmp_path, monkeypatch):
    import d02a_physical as mod

    monkeypatch.setattr(mod, "extract_linewidth_table", lambda p: [])
    monkeypatch.setattr(mod, "phase_linewidth_summary", lambda r: [])
    monkeypatch.setattr(mod, "extract_dho_spectra", lambda p: [])
    monkeypatch.setattr(mod, "extract_carrier_maps", lambda p: [])

    freeze = {
        "schema": "d02a-cspbbr3-physical-preexecution-freeze-v0.1",
        "evidence_class": "P0_D_PHYSICAL_CALIBRATION_ALREADY_VIEWED_BROAD_CONCLUSION",
    }
    contract = {
        "schema": "d02a-execution-contract-v0.1",
        "dho_license_before_execution": {
            "status": "OMEGA0_NOT_IDENTIFIABLE_FROM_LOCKED_SOURCE_TABLES",
            "chi_DHO_computed": False,
        },
    }

    result = run_d02a(
        freeze,
        contract,
        tmp_path / "f3.xlsx",
        tmp_path / "f4.xlsx",
        [],
    )
    assert result["status"] == "PHYSICAL_SOURCE_EXTRACTED_CHI_REFUSED"
    assert result["lowercase_chi"]["status"] == "REFUSED"
    assert result["lowercase_chi"]["value_records"] == []
    assert result["broader_Chi"]["master_scalar_emitted"] is False
