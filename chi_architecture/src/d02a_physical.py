from __future__ import annotations

import hashlib
import json
import math
from pathlib import Path
from typing import Any

import numpy as np
import requests
from openpyxl import load_workbook


class D02ARefusal(ValueError):
    """Fail-closed refusal for D02A physical source processing."""


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def phase_for_temperature(temp_k: float) -> str:
    t = float(temp_k)
    if not math.isfinite(t):
        raise D02ARefusal("temperature must be finite")
    if t < 360.0:
        return "ORTHORHOMBIC"
    if t < 410.0:
        return "TETRAGONAL"
    return "CUBIC"


def _candidate_urls(url: str) -> list[str]:
    mirror = url.replace(
        "https://media.springernature.com/original/springer-static/",
        "https://static-content.springer.com/",
    )
    return [url] if mirror == url else [url, mirror]


def fetch_locked_source(
    source: dict[str, Any],
    expected_sha256: str,
    destination: Path,
) -> dict[str, Any]:
    destination.parent.mkdir(parents=True, exist_ok=True)
    session = requests.Session()
    session.headers.update({"User-Agent": "Mozilla/5.0 SymC-reproducibility/1.0"})
    errors: list[str] = []

    for url in _candidate_urls(str(source["url"])):
        try:
            response = session.get(url, timeout=(15, 45))
            response.raise_for_status()
            data = response.content
            if len(data) < 1024:
                raise D02ARefusal(f"download unexpectedly small: {len(data)} bytes")
            digest = sha256_bytes(data)
            if digest != expected_sha256:
                raise D02ARefusal(
                    f"source hash mismatch for {source['id']}: {digest} != {expected_sha256}"
                )
            destination.write_bytes(data)
            return {
                "id": source["id"],
                "realized_url": url,
                "sha256": digest,
                "size_bytes": len(data),
            }
        except Exception as exc:
            errors.append(f"{url}: {exc!r}")
    raise D02ARefusal("source fetch failed: " + " | ".join(errors))


def _numeric(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return float(value)
    return None


def _complete_rows(ws, start_row: int, columns: list[str]) -> list[list[float]]:
    out: list[list[float]] = []
    for row in range(start_row, ws.max_row + 1):
        values = [_numeric(ws[f"{col}{row}"].value) for col in columns]
        if all(v is not None for v in values):
            out.append([float(v) for v in values])
    return out


def _complete_rows_between(
    ws,
    start_row: int,
    stop_before_row: int,
    columns: list[str],
) -> list[list[float]]:
    out: list[list[float]] = []
    for row in range(start_row, min(stop_before_row, ws.max_row + 1)):
        values = [_numeric(ws[f"{col}{row}"].value) for col in columns]
        if all(v is not None for v in values):
            out.append([float(v) for v in values])
    return out


def extract_linewidth_table(fig4_path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(fig4_path, data_only=True, read_only=True)
    ws = wb["4b"]
    raw = _complete_rows(ws, 3, ["A", "B", "C"])
    records = [
        {
            "temperature_K": t,
            "linewidth_meV": width,
            "linewidth_error_meV": err,
            "phase": phase_for_temperature(t),
        }
        for t, width, err in raw
    ]
    records.sort(key=lambda r: r["temperature_K"])
    for i, record in enumerate(records):
        if i == 0:
            record["change_from_previous"] = None
            continue
        prev = records[i - 1]
        delta = record["linewidth_meV"] - prev["linewidth_meV"]
        combined = math.hypot(record["linewidth_error_meV"], prev["linewidth_error_meV"])
        record["change_from_previous"] = {
            "delta_linewidth_meV": delta,
            "combined_uncertainty_meV": combined,
            "z_like_standardized_difference": None if combined <= 0.0 else delta / combined,
        }
    return records


def phase_linewidth_summary(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for phase in ["ORTHORHOMBIC", "TETRAGONAL", "CUBIC"]:
        rows = [r for r in records if r["phase"] == phase]
        if not rows:
            continue
        widths = np.asarray([r["linewidth_meV"] for r in rows], dtype=float)
        errors = np.asarray([r["linewidth_error_meV"] for r in rows], dtype=float)
        positive = errors > 0.0
        if np.any(positive):
            weights = 1.0 / np.square(errors[positive])
            weighted_mean = float(np.sum(weights * widths[positive]) / np.sum(weights))
            weighted_se = float(math.sqrt(1.0 / np.sum(weights)))
        else:
            weighted_mean = None
            weighted_se = None
        out.append(
            {
                "phase": phase,
                "row_count": len(rows),
                "temperatures_K": [r["temperature_K"] for r in rows],
                "arithmetic_mean_linewidth_meV": float(np.mean(widths)),
                "inverse_variance_weighted_mean_linewidth_meV": weighted_mean,
                "inverse_variance_weighted_standard_error_meV": weighted_se,
            }
        )
    return out


def extract_dho_spectra(fig4_path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(fig4_path, data_only=True, read_only=True)
    ws = wb["4a"]
    blocks = [
        (50.0, ["A", "B", "C"], ["D", "E"]),
        (300.0, ["F", "G", "H"], ["I", "J"]),
        (385.0, ["K", "L", "M"], ["N", "O"]),
        (500.0, ["P", "Q", "R"], ["S", "T"]),
    ]
    out: list[dict[str, Any]] = []
    for temp, data_cols, fit_cols in blocks:
        data = _complete_rows(ws, 4, data_cols)
        fit = _complete_rows(ws, 4, fit_cols)
        fit_diag: dict[str, Any] = {
            "status": "INSUFFICIENT_OVERLAP",
            "standardized_residual_rms": None,
            "comparison_point_count": 0,
        }
        if data and fit:
            data_arr = np.asarray(data, dtype=float)
            fit_arr = np.asarray(fit, dtype=float)
            order = np.argsort(fit_arr[:, 0])
            fit_x = fit_arr[order, 0]
            fit_y = fit_arr[order, 1]
            mask = (
                (data_arr[:, 0] >= fit_x[0])
                & (data_arr[:, 0] <= fit_x[-1])
                & (data_arr[:, 2] > 0.0)
            )
            if np.count_nonzero(mask) >= 3:
                predicted = np.interp(data_arr[mask, 0], fit_x, fit_y)
                residual = (data_arr[mask, 1] - predicted) / data_arr[mask, 2]
                fit_diag = {
                    "status": "DESCRIPTIVE_SOURCE_FIT_DIAGNOSTIC",
                    "standardized_residual_rms": float(np.sqrt(np.mean(np.square(residual)))),
                    "comparison_point_count": int(np.count_nonzero(mask)),
                }
        out.append(
            {
                "temperature_K": temp,
                "phase": phase_for_temperature(temp),
                "ins_data": [
                    {"energy_meV": e, "intensity": intensity, "error": err}
                    for e, intensity, err in data
                ],
                "published_dho_fit_curve": [
                    {"energy_meV": e, "intensity": intensity}
                    for e, intensity in fit
                ],
                "fit_curve_diagnostic": fit_diag,
            }
        )
    return out


def _carrier_summary(rows: list[list[float]]) -> dict[str, Any]:
    if not rows:
        return {
            "row_count": 0,
            "unique_q_count": 0,
            "q_range": None,
            "energy_range_meV": None,
            "intensity_range": None,
        }
    arr = np.asarray(rows, dtype=float)
    return {
        "row_count": int(arr.shape[0]),
        "unique_q_count": int(np.unique(arr[:, 0]).size),
        "q_range": [float(np.min(arr[:, 0])), float(np.max(arr[:, 0]))],
        "energy_range_meV": [float(np.min(arr[:, 1])), float(np.max(arr[:, 1]))],
        "intensity_range": [float(np.min(arr[:, 2])), float(np.max(arr[:, 2]))],
    }


def extract_carrier_maps(fig3_path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(fig3_path, data_only=True, read_only=True)
    specs = [
        ("Panel a", 300.0, 3, None, "SPINS_M_R"),
        ("Panel b", 385.0, 3, None, "SPINS_M_R"),
        ("Panel c G-M-R-G", 419.0, 66, 88, "CNCS_M_R_SEGMENT"),
    ]
    out: list[dict[str, Any]] = []
    for sheet, temp, start, stop, role in specs:
        ws = wb[sheet]
        rows = (
            _complete_rows(ws, start, ["A", "B", "C"])
            if stop is None
            else _complete_rows_between(ws, start, stop, ["A", "B", "C"])
        )
        out.append(
            {
                "sheet": sheet,
                "temperature_K": temp,
                "phase": phase_for_temperature(temp),
                "role": role,
                "path": "[H,1.5,0.5] M-R",
                "carrier_correspondence_status": "PATH_DEFINED_BEFORE_VALUES",
                "summary": _carrier_summary(rows),
                "rows": [
                    {"q_rlu": q, "energy_meV": e, "intensity": intensity}
                    for q, e, intensity in rows
                ],
            }
        )
    return out


def run_d02a(
    freeze: dict[str, Any],
    execution_contract: dict[str, Any],
    fig3_path: Path,
    fig4_path: Path,
    source_records: list[dict[str, Any]],
) -> dict[str, Any]:
    if freeze.get("schema") != "d02a-cspbbr3-physical-preexecution-freeze-v0.1":
        raise D02ARefusal("unexpected D02A freeze")
    if execution_contract.get("schema") != "d02a-execution-contract-v0.1":
        raise D02ARefusal("unexpected D02A execution contract")

    linewidth = extract_linewidth_table(fig4_path)
    spectra = extract_dho_spectra(fig4_path)
    carriers = extract_carrier_maps(fig3_path)

    dho_license = dict(execution_contract["dho_license_before_execution"])
    if dho_license.get("chi_DHO_computed") is not False:
        raise D02ARefusal("D02A v0.1 must not compute chi under the frozen source license")

    return {
        "schema": "d02a-cspbbr3-physical-result-v0.1",
        "status": "PHYSICAL_SOURCE_EXTRACTED_CHI_REFUSED",
        "evidence_class": freeze["evidence_class"],
        "confirmatory_credit": False,
        "source_records": source_records,
        "linewidth_records": linewidth,
        "phase_linewidth_summary": phase_linewidth_summary(linewidth),
        "dho_spectra": spectra,
        "carrier_maps": carriers,
        "dho_license": dho_license,
        "lowercase_chi": {
            "status": "REFUSED",
            "value_records": [],
            "reason": dho_license["status"],
        },
        "broader_Chi": {
            "status": "NATIVE_CARRIER_ARCHITECTURE_RETAINED",
            "master_scalar_emitted": False,
            "components": [
                "temperature",
                "structural_phase",
                "M-R carrier path",
                "source DHO-derived linewidth",
                "source spectral response",
            ],
        },
        "interpretation_ceiling": (
            "P0-D physical calibration only. Broad qualitative behavior was known before freeze; "
            "no untouched confirmation or recovery/resilience claim."
        ),
    }


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))
