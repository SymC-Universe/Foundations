from __future__ import annotations

import json
import math
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from d02a_physical import (
    D02ARefusal,
    extract_linewidth_table,
    phase_linewidth_summary,
)


def _numeric_matrix(ws, cell_range: str) -> np.ndarray:
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    values: list[list[float]] = []
    for row in ws.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
        values_only=True,
    ):
        numeric_row: list[float] = []
        for value in row:
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                raise D02ARefusal(f"non-numeric cell inside frozen range {cell_range}: {value!r}")
            x = float(value)
            if not math.isfinite(x):
                raise D02ARefusal(f"non-finite cell inside frozen range {cell_range}")
            numeric_row.append(x)
        values.append(numeric_row)
    arr = np.asarray(values, dtype=float)
    if arr.ndim != 2:
        raise D02ARefusal(f"unexpected matrix rank for {cell_range}")
    return arr


def _numeric_vector(ws, cell_range: str) -> np.ndarray:
    arr = _numeric_matrix(ws, cell_range)
    if 1 not in arr.shape:
        raise D02ARefusal(f"frozen vector range is not one-dimensional: {cell_range}")
    return arr.reshape(-1)


def load_frozen_carrier_matrix(
    fig3_path: Path,
    geometry: dict[str, Any],
) -> dict[str, Any]:
    wb = load_workbook(fig3_path, data_only=True, read_only=True)
    ws = wb[str(geometry["sheet"])]

    q = _numeric_vector(ws, str(geometry["q_axis_range"]))
    energy = _numeric_vector(ws, str(geometry["energy_axis_range"]))
    intensity = _numeric_matrix(ws, str(geometry["intensity_matrix_range"]))

    expected_shape = tuple(int(x) for x in geometry["matrix_shape_energy_by_q"])
    if intensity.shape != expected_shape:
        raise D02ARefusal(
            f"carrier matrix shape mismatch: {intensity.shape} != {expected_shape}"
        )
    if q.size != expected_shape[1]:
        raise D02ARefusal("q axis length does not match carrier matrix")
    if energy.size != expected_shape[0]:
        raise D02ARefusal("energy axis length does not match carrier matrix")
    if np.any(np.diff(q) <= 0.0):
        raise D02ARefusal("q axis must be strictly increasing")
    if np.any(np.diff(energy) <= 0.0):
        raise D02ARefusal("energy axis must be strictly increasing")

    return {
        "q": q,
        "energy": energy,
        "intensity": intensity,
    }


def _comparison_grid(contract: dict[str, Any]) -> np.ndarray:
    spec = contract["cross_instrument_comparison"]["comparison_grid_meV"]
    start = float(spec["start"])
    stop = float(spec["stop"])
    step = float(spec["step"])
    n = int(round((stop - start) / step))
    grid = start + step * np.arange(n + 1, dtype=float)
    if abs(float(grid[-1]) - stop) > 1e-10:
        raise D02ARefusal("comparison grid does not land on frozen stop")
    return grid


def normalize_carrier_shape(
    q: np.ndarray,
    energy: np.ndarray,
    intensity: np.ndarray,
    common_energy: np.ndarray,
) -> dict[str, Any]:
    q = np.asarray(q, dtype=float)
    energy = np.asarray(energy, dtype=float)
    intensity = np.asarray(intensity, dtype=float)
    common_energy = np.asarray(common_energy, dtype=float)

    if intensity.shape != (energy.size, q.size):
        raise D02ARefusal("intensity matrix does not match energy/q axes")
    if common_energy[0] < energy[0] - 1e-12 or common_energy[-1] > energy[-1] + 1e-12:
        raise D02ARefusal("common energy grid falls outside source support")

    normalized = np.zeros((common_energy.size, q.size), dtype=float)
    metrics: list[dict[str, Any]] = []

    for j, q_value in enumerate(q):
        interpolated = np.interp(common_energy, energy, intensity[:, j])
        positive = np.clip(interpolated, 0.0, None)
        total = float(np.sum(positive))
        if not math.isfinite(total) or total <= 0.0:
            normalized[:, j] = np.nan
            metrics.append({
                "q_rlu": float(q_value),
                "status": "NON_IDENTIFIABLE",
                "normalized_energy_centroid_meV": None,
                "normalized_energy_rms_width_meV": None,
                "peak_energy_meV_on_common_grid": None,
            })
            continue

        weights = positive / total
        normalized[:, j] = weights
        centroid = float(np.sum(weights * common_energy))
        width = float(np.sqrt(np.sum(weights * np.square(common_energy - centroid))))
        peak = float(common_energy[int(np.argmax(weights))])

        metrics.append({
            "q_rlu": float(q_value),
            "status": "IDENTIFIABLE",
            "normalized_energy_centroid_meV": centroid,
            "normalized_energy_rms_width_meV": width,
            "peak_energy_meV_on_common_grid": peak,
        })

    return {
        "normalized_intensity": normalized,
        "q_metrics": metrics,
    }


def _paired_metric_arrays(
    left: list[dict[str, Any]],
    right: list[dict[str, Any]],
    field: str,
) -> tuple[np.ndarray, np.ndarray]:
    if len(left) != len(right):
        raise D02ARefusal("paired carrier metrics have different lengths")
    a: list[float] = []
    b: list[float] = []
    for lrow, rrow in zip(left, right):
        if abs(float(lrow["q_rlu"]) - float(rrow["q_rlu"])) > 1e-12:
            raise D02ARefusal("paired q grids do not match")
        lv = lrow[field]
        rv = rrow[field]
        if lv is None or rv is None:
            continue
        a.append(float(lv))
        b.append(float(rv))
    return np.asarray(a, dtype=float), np.asarray(b, dtype=float)


def _rms_difference(a: np.ndarray, b: np.ndarray) -> float | None:
    if a.size == 0:
        return None
    return float(np.sqrt(np.mean(np.square(a - b))))


def _pearson(a: np.ndarray, b: np.ndarray) -> float | None:
    if a.size < 2:
        return None
    if float(np.std(a)) <= 1e-15 or float(np.std(b)) <= 1e-15:
        return None
    return float(np.corrcoef(a, b)[0, 1])


def pairwise_carrier_metrics(
    left: dict[str, Any],
    right: dict[str, Any],
) -> dict[str, Any]:
    lmetrics = left["q_metrics"]
    rmetrics = right["q_metrics"]

    c1, c2 = _paired_metric_arrays(
        lmetrics, rmetrics, "normalized_energy_centroid_meV"
    )
    w1, w2 = _paired_metric_arrays(
        lmetrics, rmetrics, "normalized_energy_rms_width_meV"
    )
    p1, p2 = _paired_metric_arrays(
        lmetrics, rmetrics, "peak_energy_meV_on_common_grid"
    )

    return {
        "left_temperature_K": left["temperature_K"],
        "right_temperature_K": right["temperature_K"],
        "paired_identifiable_q_count_centroid": int(c1.size),
        "centroid_rms_difference_meV": _rms_difference(c1, c2),
        "width_rms_difference_meV": _rms_difference(w1, w2),
        "peak_energy_rms_difference_meV": _rms_difference(p1, p2),
        "centroid_pearson_r": _pearson(c1, c2),
        "width_pearson_r": _pearson(w1, w2),
    }


def run_d02a_v02(
    freeze: dict[str, Any],
    source_lock: dict[str, Any],
    v01_execution_contract: dict[str, Any],
    v02_contract: dict[str, Any],
    fig3_path: Path,
    fig4_path: Path,
    source_records: list[dict[str, Any]],
) -> dict[str, Any]:
    if v02_contract.get("schema") != "d02a-v0.2-carrier-parser-contract":
        raise D02ARefusal("unexpected D02A v0.2 parser contract")
    if freeze.get("schema") != "d02a-cspbbr3-physical-preexecution-freeze-v0.1":
        raise D02ARefusal("unexpected D02A scientific freeze")

    linewidth = extract_linewidth_table(fig4_path)
    linewidth_summary = phase_linewidth_summary(linewidth)
    common_energy = _comparison_grid(v02_contract)

    geometry = v02_contract["matrix_geometry"]
    ordered = [
        ("300K_ORTHORHOMBIC", 300.0, "ORTHORHOMBIC"),
        ("385K_TETRAGONAL", 385.0, "TETRAGONAL"),
        ("419K_CUBIC", 419.0, "CUBIC"),
    ]

    carriers: list[dict[str, Any]] = []
    reference_q: np.ndarray | None = None

    for key, temperature, phase in ordered:
        raw = load_frozen_carrier_matrix(fig3_path, geometry[key])
        if reference_q is None:
            reference_q = raw["q"]
        elif not np.allclose(raw["q"], reference_q, atol=1e-12, rtol=0.0):
            raise D02ARefusal("frozen source q grids are not identical")

        shape = normalize_carrier_shape(
            raw["q"],
            raw["energy"],
            raw["intensity"],
            common_energy,
        )
        carriers.append({
            "key": key,
            "temperature_K": temperature,
            "phase": phase,
            "sheet": geometry[key]["sheet"],
            "q_axis_rlu": [float(x) for x in raw["q"]],
            "energy_axis_meV": [float(x) for x in raw["energy"]],
            "raw_intensity_matrix": raw["intensity"].tolist(),
            "common_energy_grid_meV": [float(x) for x in common_energy],
            "normalized_intensity_matrix": shape["normalized_intensity"].tolist(),
            "q_metrics": shape["q_metrics"],
            "absolute_intensity_comparison_allowed": False,
        })

    pairwise = [
        pairwise_carrier_metrics(carriers[0], carriers[1]),
        pairwise_carrier_metrics(carriers[1], carriers[2]),
        pairwise_carrier_metrics(carriers[0], carriers[2]),
    ]

    dho_license = dict(v01_execution_contract["dho_license_before_execution"])
    if dho_license.get("chi_DHO_computed") is not False:
        raise D02ARefusal("D02A v0.2 must preserve lowercase chi refusal")

    return {
        "schema": "d02a-cspbbr3-physical-result-v0.2",
        "status": "PHYSICAL_CARRIER_MATRIX_RECONSTRUCTED_CHI_REFUSED_V0.2",
        "evidence_class": "P0_D_POST_RESULT_PARSER_DEVELOPMENT_ONLY",
        "confirmatory_credit": False,
        "parent_v0_1_status": "PARTIAL_VALID_LINEWIDTH__CARRIER_PARSER_INADEQUATE__CHI_REFUSED",
        "source_records": source_records,
        "source_hash_lock": {
            "FIG3": source_lock["sources"]["FIG3"]["sha256"],
            "FIG4": source_lock["sources"]["FIG4"]["sha256"],
        },
        "linewidth_records": linewidth,
        "phase_linewidth_summary": linewidth_summary,
        "carrier_maps": carriers,
        "pairwise_carrier_shape_metrics": pairwise,
        "dho_license": dho_license,
        "lowercase_chi": {
            "status": "REFUSED",
            "value_records": [],
            "reason": dho_license["status"],
        },
        "broader_Chi": {
            "status": "NATIVE_CARRIER_ARCHITECTURE_RETAINED_NO_MASTER_SCALAR",
            "master_scalar_emitted": False,
            "carrier": "M-R reciprocal-space sector",
            "organization_metrics": [
                "q-resolved normalized energy centroid",
                "q-resolved normalized RMS width",
                "q-resolved peak energy on common grid",
            ],
        },
        "comparison_firewall": {
            "absolute_intensity_compared_across_instruments": False,
            "categorical_similarity_threshold_used": False,
            "DHO_parameters_fitted_from_carrier_maps": False,
        },
        "interpretation_ceiling": (
            "P0-D post-result physical parser development only. v0.2 corrects the "
            "source matrix representation but cannot become untouched confirmation."
        ),
    }


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))
