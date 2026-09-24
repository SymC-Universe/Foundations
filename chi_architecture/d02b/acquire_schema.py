from __future__ import annotations

import hashlib
import json
from pathlib import Path

import openpyxl
import requests


ROOT = Path(__file__).resolve().parents[1]
CACHE = ROOT / "results" / "d02b_sources"
CACHE.mkdir(parents=True, exist_ok=True)

SOURCE_URL = "https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41467-025-59292-2/MediaObjects/41467_2025_59292_MOESM3_ESM.xlsx"
CODE_URL = "https://zenodo.org/records/15162385/files/Figure2simulation.py?download=1"


def download(url: str, path: Path) -> None:
    headers = {"User-Agent": "SymC-reproducibility/1.0 (+https://github.com/SymC-Universe/Foundations)"}
    with requests.get(url, headers=headers, stream=True, timeout=180) as r:
        r.raise_for_status()
        with path.open("wb") as f:
            for chunk in r.iter_content(chunk_size=1024 * 1024):
                if chunk:
                    f.write(chunk)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def safe_schema(workbook_path: Path) -> dict:
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=False)
    sheets = []
    for ws in wb.worksheets:
        string_labels = []
        type_counts = {"string": 0, "number": 0, "formula": 0, "boolean": 0, "date": 0, "other": 0}
        scanned_nonempty = 0
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if v is None:
                    continue
                scanned_nonempty += 1
                dt = cell.data_type
                if dt == "s":
                    type_counts["string"] += 1
                    if len(string_labels) < 30:
                        text = str(v).strip()
                        if text and len(text) <= 160:
                            string_labels.append({"cell": cell.coordinate, "text": text})
                elif dt == "n":
                    type_counts["number"] += 1
                elif dt == "f":
                    type_counts["formula"] += 1
                elif dt == "b":
                    type_counts["boolean"] += 1
                elif dt == "d":
                    type_counts["date"] += 1
                else:
                    type_counts["other"] += 1
        sheets.append({
            "title": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "nonempty_cell_count": scanned_nonempty,
            "cell_type_counts": type_counts,
            "string_labels_only": string_labels,
        })
    wb.close()
    return {"sheet_count": len(sheets), "sheets": sheets}


def main() -> None:
    workbook = CACHE / "41467_2025_59292_MOESM3_ESM.xlsx"
    code = CACHE / "Figure2simulation.py"
    if not workbook.exists():
        download(SOURCE_URL, workbook)
    if not code.exists():
        download(CODE_URL, code)

    out = {
        "schema": "d02b-source-schema-v0.1",
        "status": "SCHEMA_ONLY_NO_NUMERIC_CELL_VALUES_EMITTED",
        "sources": {
            "workbook": {
                "url": SOURCE_URL,
                "filename": workbook.name,
                "size_bytes": workbook.stat().st_size,
                "sha256": sha256(workbook),
            },
            "author_code": {
                "url": CODE_URL,
                "filename": code.name,
                "size_bytes": code.stat().st_size,
                "sha256": sha256(code),
            },
        },
        "workbook_schema": safe_schema(workbook),
        "guard": "Numeric cell values are deliberately omitted from this checkpoint artifact.",
    }
    output = ROOT / "results" / "D02B_SOURCE_SCHEMA_v0.1.json"
    output.write_text(json.dumps(out, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps({
        "status": out["status"],
        "workbook_sha256": out["sources"]["workbook"]["sha256"],
        "code_sha256": out["sources"]["author_code"]["sha256"],
        "sheet_count": out["workbook_schema"]["sheet_count"],
        "output": str(output),
    }, indent=2))


if __name__ == "__main__":
    main()
