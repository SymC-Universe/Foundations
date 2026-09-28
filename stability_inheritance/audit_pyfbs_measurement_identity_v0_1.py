#!/usr/bin/env python3
import hashlib, json, urllib.request
from pathlib import Path
import numpy as np
from openpyxl import load_workbook

BASE="https://gitlab.com/pyFBS/pyFBS_data/-/raw/master/lab_testbench/Measurements/"
FRF=["Y_A.p","Y_B.p","Y_AB.p"]
META=["coupling_example.xlsx","decoupling_example.xlsx","AM_Measurements.xlsx"]
OUT=Path("stability_inheritance/results/pyfbs_identity_audit")
RAW=OUT/"raw"
OUT.mkdir(parents=True,exist_ok=True); RAW.mkdir(parents=True,exist_ok=True)
rep={"status":"SHAPE_METADATA_IDENTITY_ONLY_NO_FRF_VALUES_SCORED","frf":[],"metadata":[]}

for name in FRF:
    p=RAW/name; urllib.request.urlretrieve(BASE+name,p)
    b=p.read_bytes()
    obj=np.load(p,allow_pickle=True)
    freq=np.asarray(obj[0])
    arr=np.asarray(obj[1])
    rep["frf"].append({
      "name":name,"size_bytes":len(b),"sha256":hashlib.sha256(b).hexdigest(),
      "freq_shape":list(freq.shape),"frf_shape":list(arr.shape),"frf_ndim":int(arr.ndim),
      "dtype":str(arr.dtype)
    })

for name in META:
    p=RAW/name; urllib.request.urlretrieve(BASE+name,p)
    b=p.read_bytes()
    wb=load_workbook(p,data_only=True,read_only=True)
    sheets=[]
    for ws in wb.worksheets:
        sheets.append({"name":ws.title,"rows":ws.max_row,"columns":ws.max_column})
    rep["metadata"].append({"name":name,"size_bytes":len(b),"sha256":hashlib.sha256(b).hexdigest(),"sheets":sheets})

(OUT/"identity.json").write_text(json.dumps(rep,indent=2))
print(json.dumps(rep,indent=2))
for p in RAW.iterdir(): p.unlink()
RAW.rmdir()
