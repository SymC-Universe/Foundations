#!/usr/bin/env python3
import json, urllib.request
from pathlib import Path
from openpyxl import load_workbook

URL="https://gitlab.com/pyFBS/pyFBS_data/-/raw/master/lab_testbench/Measurements/coupling_example.xlsx"
OUT=Path("stability_inheritance/results/pyfbs_lab_metadata")
OUT.mkdir(parents=True,exist_ok=True)
path=OUT/"coupling_example.xlsx"
urllib.request.urlretrieve(URL,path)

wb=load_workbook(path,data_only=True,read_only=True)
report={"status":"METADATA_ONLY_NO_FRF_VALUES","source_url":URL,"sheets":[]}
for ws in wb.worksheets:
    rows=list(ws.iter_rows(values_only=True))
    header=list(rows[0]) if rows else []
    report["sheets"].append({
        "name":ws.title,
        "max_row":ws.max_row,
        "max_column":ws.max_column,
        "header":[str(x) if x is not None else None for x in header],
        "rows":[[x for x in row] for row in rows[1:]] if ws.title in ("VP_Channels","VP_RefChannels") else None
    })
(OUT/"metadata.json").write_text(json.dumps(report,indent=2,default=str))
path.unlink()
print(json.dumps(report,indent=2,default=str))
